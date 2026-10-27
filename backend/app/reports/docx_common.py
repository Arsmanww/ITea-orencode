from io import BytesIO

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.table import Table, _Cell

FONT = "Times New Roman"
ORGANIZATION = "Региональное отделение Федерации спортивного программирования Оренбургской области"
LEVEL_LABELS = {"regional": "региональное", "interregional": "межрегиональное", "all_russian": "всероссийское"}
PLAN_LABELS = {
    "ekp_minsport": "Единый календарный план Минспорта России",
    "oren_regional": "Календарный план Оренбургской области",
    "none": "не включено",
}
FORMAT_LABELS = {"online": "онлайн", "offline": "очно"}
STATUS_LABELS = {"ranked": "", "did_not_start": "не явился", "disqualified": "дисквалифицирован"}


def new_document(landscape: bool = False) -> Document:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name, style.font.size = FONT, Pt(12)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    style.paragraph_format.space_after = Pt(4)
    section = doc.sections[0]
    if landscape:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = section.bottom_margin = Cm(1.8)
    section.left_margin, section.right_margin = Cm(2.2), Cm(1.5)
    return doc


def para(doc: Document, text: str = "", *, bold: bool = False, italic: bool = False, size: float | None = None,
         align: str = "left", after: float | None = None, before: float | None = None, color: str | None = None):
    p = doc.add_paragraph()
    p.alignment = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER,
                   "right": WD_ALIGN_PARAGRAPH.RIGHT, "justify": WD_ALIGN_PARAGRAPH.JUSTIFY}[align]
    if after is not None:
        p.paragraph_format.space_after = Pt(after)
    if before is not None:
        p.paragraph_format.space_before = Pt(before)
    if text:
        run = p.add_run(text)
        run.bold, run.italic = bold, italic
        if size:
            run.font.size = Pt(size)
        if color:
            run.font.color.rgb = RGBColor.from_string(color)
    return p


def shade(cell: _Cell, fill: str) -> None:
    props = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    props.append(shd)


def repeat_header(row) -> None:
    props = row._tr.get_or_add_trPr()
    flag = OxmlElement("w:tblHeader")
    flag.set(qn("w:val"), "true")
    props.append(flag)


def keep_row_together(row) -> None:
    props = row._tr.get_or_add_trPr()
    props.append(OxmlElement("w:cantSplit"))


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths_cm: list[float], *, font_size: float = 10,
              header_fill: str = "D9E2F3", center_cols: set[int] | None = None) -> Table:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    center_cols = center_cols or set()

    def fill(cell: _Cell, text: str, bold: bool, centered: bool) -> None:
        cell.text = ""
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if centered else WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run(text)
        run.bold, run.font.size = bold, Pt(font_size)

    for i, h in enumerate(headers):
        fill(table.rows[0].cells[i], h, True, True)
        shade(table.rows[0].cells[i], header_fill)
    repeat_header(table.rows[0])
    for values in rows:
        row = table.add_row()
        keep_row_together(row)
        for i, v in enumerate(values):
            fill(row.cells[i], v, False, i in center_cols)
    for row in table.rows:
        for i, w in enumerate(widths_cm):
            row.cells[i].width = Cm(w)
    return table


def signature_block(doc: Document, role: str, name: str = "") -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(14), Pt(2)
    p.add_run(f"{role}").bold = True
    p.add_run(f"\t\t______________  /  {name or '________________________'}  /")
    p.paragraph_format.tab_stops.add_tab_stop(Cm(7.5))
    p.paragraph_format.tab_stops.add_tab_stop(Cm(9.0))


def footer_with_page_number(doc: Document, text: str) -> None:
    footer = doc.sections[0].footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(f"{text}  ·  стр. ")
    run.font.size = Pt(9)
    field_run = p.add_run()
    field_run.font.size = Pt(9)
    for kind, instr in (("begin", None), (None, " PAGE "), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = instr
        field_run._r.append(el)


def fmt_time(seconds) -> str:
    if seconds is None:
        return ""
    total = float(seconds)
    h, rest = divmod(int(total), 3600)
    m, s = divmod(rest, 60)
    frac = total - int(total)
    tail = f".{int(round(frac * 1000)):03d}" if frac > 0.0005 else ""
    return f"{h}:{m:02d}:{s:02d}{tail}" if h else f"{m}:{s:02d}{tail}"


def fmt_date(value) -> str:
    return value.strftime("%d.%m.%Y") if value else ""


def fmt_num(value) -> str:
    if value is None:
        return ""
    text = f"{float(value):.2f}".rstrip("0").rstrip(".")
    return text.replace(".", ",")


def to_bytes(doc: Document) -> bytes:
    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
